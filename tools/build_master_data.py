#!/usr/bin/env python3
"""
Rebuild src/master-data.json from the four ETT master workbooks.

Usage:
    pip install openpyxl
    python3 build_master_data.py  <folder containing the 4 xlsx>  [-o ../src/master-data.json]

Expected files in that folder (name match is case-insensitive, substring):
    SOFTNESS.xlsx            softness standard per article / thickness / zone
    Wet blue standard.xlsx   wet-blue selection standard  (sheet "New selection")
    COLOUR CODE.xlsx         valid colour codes per article
    REJECT REASON.xlsx       reject / defect reason list

Output JSON shape:
{
  "ART": [ {a: article, t: "1.3-1.5", b: brand, s: [neck, belly, butt]} ],
  "WB" : [ {a,t,n, ty: tannery, bd: brandDesc, g: catalogueGrade,
            h: hides,  sp: supplier,  mc: materialCode,       # main option
            ah,asp,amc,   ch,csp,cmc,   nh,nsp,nmc} ],        # alt / calf / new options
  "COL": { article: [[code, description, isActive0or1], ...] },
  "RJ" : [ [codeNo, shortCode, nameEn, nameTh, groupIdx, phase, [stages]] ],
  "GRP": ["", "MANUFACTURING", "NATURAL DEFECTS & MATERIAL", "TRANSFER", "COLOR", "OTHER"],
  "FIN": ["", "NUBUCK LIGHT", "SEMI ANILINE", "NUBUCK DARK", "PIGMENTED",
          "BUFFED-OIL", "EMBOSS", "CORRECTED GRAIN"]
}
"""
import sys, os, re, json, glob, argparse

try:
    import openpyxl
except ImportError:
    sys.exit("pip install openpyxl")

STAGES = ["WBTO", "RTO", "CTO", "FHO", "GIP"]
GROUP_KEYS = ["", "1", "2", "3", "5", "9"]   # index = groupIdx used in the app
GROUP_NAMES = ["", "MANUFACTURING", "NATURAL DEFECTS & MATERIAL", "TRANSFER", "COLOR", "OTHER"]
FINISH_GRADES = ["", "NUBUCK LIGHT", "SEMI ANILINE", "NUBUCK DARK", "PIGMENTED",
                 "BUFFED-OIL", "EMBOSS", "CORRECTED GRAIN"]

U = lambda v: (str(v).strip().upper() if v not in (None, "") else "")
dec = lambda s: s.replace(",", ".")          # 1,3-1,5 -> 1.3-1.5


def find(folder, needle):
    for p in glob.glob(os.path.join(folder, "*.xlsx")):
        if needle.lower().replace(" ", "") in os.path.basename(p).lower().replace(" ", "").replace("_", ""):
            return p
    sys.exit("missing workbook matching %r in %s" % (needle, folder))


# ---------------------------------------------------------------- softness
def load_softness(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    # Every row is kept. 131 of the 509 article+thickness keys have more than one row:
    # 33 differ by brand (a genuinely different target per customer) and the rest are
    # duplicate/conflicting entries in the master. The app picks the row whose brand
    # matches the lot's customer, and averages when several still match — so conflicts
    # stay visible instead of being silently resolved here.
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        a = U(r[2])
        if not a:
            continue
        t = dec(U(r[3])).replace(" MM", "").strip()
        f = lambda v: (round(float(v), 2) if isinstance(v, (int, float)) else None)
        out.append({"a": a, "t": t, "b": U(r[0]), "s": [f(r[4]), f(r[5]), f(r[6])]})
    return out


# ---------------------------------------------------------------- wet blue
NAME_RE = re.compile(r"^(.*?)\s+(\d,\d-\d,\d)\s*MM(.*)$")

def load_wetblue(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["New selection"] if "New selection" in wb.sheetnames else wb.worksheets[0]
    out = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        if not r[1]:
            continue
        m = NAME_RE.match(str(r[1]).strip())
        if m:
            a, t, suffix = U(m.group(1)), dec(m.group(2)), U(m.group(3)).strip(" ()")
        else:
            a, t, suffix = U(r[1]), "", ""
        row = {"a": a, "t": t, "n": suffix,
               "ty": U(r[3]), "bd": U(r[2]), "g": U(r[4]),
               "h": U(r[5]),  "sp": U(r[6]),  "mc": U(r[7]),     # catalogue option
               "ah": U(r[9]), "asp": U(r[10]), "amc": U(r[11]),  # alternative
               "ch": U(r[13]), "csp": U(r[14]), "cmc": U(r[15]), # calf option
               "nh": U(r[17]), "nsp": U(r[18]), "nmc": U(r[19])} # new option
        for k in ("g", "h", "sp", "mc", "ah", "asp", "amc", "ch", "csp", "cmc", "nh", "nsp", "nmc"):
            if row[k] in ("NONE", "0-0"):
                row[k] = ""
        out.append(row)
    return out


# ---------------------------------------------------------------- colour
def load_colour(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        validity, article, code, desc = U(r[0]), U(r[2]), U(r[4]), U(r[5])
        if not article or not code:
            continue
        active = 1 if validity.startswith("ACTIVE") and "CANCEL" not in validity else 0
        out.setdefault(article, []).append([code, desc, active])
    return out


# ---------------------------------------------------------------- reject reasons
#
# Which stage should offer which reason. The workbook's Phase column is the base;
# everything else is a rule so the mapping can be re-derived when the list changes.
# Edit these rules — do not hand-edit the JSON.
#
FINISHING_OP = re.compile(r"BRUSH|COVERAGE|EFFECT|SPRAY|LACQUER|COAT|EMBOSS|FOIL|MILL|BUFF|"
                          r"TIPPING|POLISH|PLATE|PRINT|PEBBLE|TAPE|ADHES|PEEL")

def stages_for(name, group_idx, phase):
    s = set()
    if phase == "WET BLUE":
        s |= {"WBTO", "RTO"}
    if phase == "CRUST/FINISHED":
        s |= {"CTO", "FHO"}
    if phase == "FINISHED":
        s |= {"FHO"}
    if group_idx == 2:                                   # natural / material defect
        s |= {"WBTO", "RTO", "CTO", "FHO"}
    if re.search(r"SUBSTANCE|THICK|SHAV|SPLIT|TRIM", name):
        s |= {"RTO", "CTO"}
    if re.search(r"COLOR|COLOUR|DYE|SHADE|TONE|BRONZ", name) or group_idx == 5:
        s |= {"RTO", "CTO"}
    if re.search(r"MOULD|BACTERIA|FUNG|LIME|CHROME|KNIFE|FLESH|SALT|WETBLUE|WET BLUE|"
                 r"RED SPOT|BLUE SPOT|FIBER|FIBRE", name):
        s |= {"WBTO", "RTO"}
    if re.search(r"HARD|SOFT|LOOSE|BREAK|MILL|DRY|OIL|WAX|GREASE", name):
        s |= {"CTO", "FHO"}
    if FINISHING_OP.search(name):                        # finishing operation, not a wet-end defect
        s -= {"WBTO", "RTO"}
        s |= {"FHO"}
    s |= {"GIP"}                                         # final grading can reject for anything
    return [x for x in STAGES if x in s]


def load_reject(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[2]:
            continue
        g = U(r[4])
        gi = GROUP_KEYS.index(g.split(" ")[0]) if g and g.split(" ")[0] in GROUP_KEYS else 0
        name, phase = U(r[2]), U(r[5])
        out.append([str(r[0]).strip(), U(r[1]), name,
                    (str(r[3]).strip() if r[3] else ""), gi, phase,
                    stages_for(name, gi, phase)])
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("folder")
    ap.add_argument("-o", "--out", default="master-data.json")
    ap.add_argument("--csv", help="also write the reject-reason stage map to this CSV for review")
    a = ap.parse_args()

    data = {
        "ART": load_softness(find(a.folder, "softness")),
        "WB":  load_wetblue(find(a.folder, "wetbluestandard")),
        "COL": load_colour(find(a.folder, "colourcode")),
        "RJ":  load_reject(find(a.folder, "rejectreason")),
        "GRP": GROUP_NAMES,
        "FIN": FINISH_GRADES,
    }
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(data, f, separators=(",", ":"), ensure_ascii=False)

    print("articles      %5d" % len({x["a"] for x in data["ART"]} | {x["a"] for x in data["WB"]}))
    print("softness rows %5d" % len(data["ART"]))
    print("wet-blue rows %5d" % len(data["WB"]))
    print("colour codes  %5d" % sum(len(v) for v in data["COL"].values()))
    print("reject reasons%5d" % len(data["RJ"]))
    for s in STAGES:
        print("  %-5s offers %3d reasons" % (s, sum(1 for x in data["RJ"] if s in x[6])))
    print("wrote", a.out)

    if a.csv:
        import csv
        with open(a.csv, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["Code No.", "Code", "Description", "TH Description",
                        "Group", "Phase"] + STAGES)
            for r in data["RJ"]:
                w.writerow([r[0], r[1], r[2], r[3], GROUP_NAMES[r[4]], r[5]] +
                           ["Y" if s in r[6] else "" for s in STAGES])
        print("wrote", a.csv)


if __name__ == "__main__":
    main()
